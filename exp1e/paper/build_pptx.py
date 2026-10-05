"""원고 v3 그림 파일: 한 그림을 한 슬라이드에, 본문 순서대로(그래픽 초록 → 그림 1–6). 보충 그림 S1은 보충 Word 파일에 있다.
그림은 figures/*.png(300 dpi)를 슬라이드 안에 비율을 유지한 채 최대 크기로 넣는다.
출력: paper/manuscript_v3_figures.pptx
"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

import sys
V4 = sys.argv[1:] == ["v4"]  # v4: 그래픽 초록은 별도 Word 파일이므로 그림 1–6만
P = Path(__file__).resolve().parent; FIG = P / "figures"; OUT = P / "v4" / "Figures_manuscript_v4.pptx" if V4 else P / "manuscript_v3_figures.pptx"
ITEMS = [("Graphical abstract", "visual_abstract.png"), ("Figure 1", "fig1.png"), ("Figure 2", "fig2.png"), ("Figure 3", "fig3.png"),
         ("Figure 4", "fig4.png"), ("Figure 5", "fig5.png"), ("Figure 6", "fig6.png")]
if V4: ITEMS = ITEMS[1:]
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height; M = Inches(0.3); TOP = Inches(0.75)
for label, fn in ITEMS:
    s = prs.slides.add_slide(prs.slide_layouts[6])  # 빈 슬라이드
    tb = s.shapes.add_textbox(M, Inches(0.15), SW - 2 * M, Inches(0.5)).text_frame; tb.text = label
    r = tb.paragraphs[0].runs[0]; r.font.size = Pt(20); r.font.bold = True; r.font.name = "Arial"; r.font.color.rgb = RGBColor(0, 0, 0)
    w, h = Image.open(FIG / fn).size; bw, bh = SW - 2 * M, SH - TOP - M
    sc = min(bw / w, bh / h); pw, ph = int(w * sc), int(h * sc)
    s.shapes.add_picture(str(FIG / fn), int((SW - pw) / 2), int(TOP + (bh - ph) / 2), pw, ph)
prs.save(OUT); print(f"저장: {OUT.name} | 슬라이드 {len(ITEMS)}")
